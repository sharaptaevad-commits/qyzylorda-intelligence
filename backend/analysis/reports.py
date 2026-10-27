# -*- coding: utf-8 -*-
"""Report export: PDF (reportlab) and Excel (openpyxl)."""
import io
import os
import time
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# register a font that supports Cyrillic (DejaVu often available; fallback Helvetica)
_FONT = "Helvetica"
for cand in [
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\calibri.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]:
    if os.path.exists(cand):
        try:
            pdfmetrics.registerFont(TTFont("Cyr", cand))
            _FONT = "Cyr"
            break
        except Exception:
            pass

def _rows_for_table(items, headers):
    return [headers] + [[_cell(v) for v in item] for item in items]

def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".")
    return str(v)

def export_pdf(doc, comp, summary, findings, out_path):
    doc_title = doc.get("filename", "document")
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(out_path, pagesize=landscape(A4),
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    H = styles["Heading2"].fontName
    story = []
    story.append(Paragraph("Document Intelligence — Report", styles["Title"]))
    story.append(Paragraph(f"File: {doc_title} | {time.strftime('%Y-%m-%d %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 6))

    # summary
    srows = [[k, v] for k, v in summary.items() if k != "sources"]
    st = Table(srows, colWidths=[45 * mm, 345 * mm])
    st.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), _FONT),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2ff")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(st)
    story.append(Spacer(1, 10))

    def section(title, items, headers):
        if not items:
            story.append(Paragraph(title, styles["Heading3"]))
            story.append(Paragraph("—", styles["Normal"]))
            story.append(Spacer(1, 6))
            return
        data = _rows_for_table(items, headers)
        t = Table(data, colWidths=[360 * mm / len(headers)] * len(headers), repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), _FONT),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        story.append(Paragraph(title, styles["Heading3"]))
        story.append(t)
        story.append(Spacer(1, 8))

    section("A. Matching values", comp["matched"],
            ["metric", "region", "period", "doc_value", "site_value", "pct_diff"])
    section("B. Discrepancies", comp["discrepancies"],
            ["metric", "region", "period", "doc_value", "site_value", "pct_diff", "possible_reasons"])
    section("C. New in document", comp["new_in_document"],
            ["metric", "region", "period", "doc_value", "doc_source"])
    section("D. Not in document", comp["missing_in_document"][:20],
            ["metric", "region", "period", "site_value", "site_source"])

    story.append(Paragraph("Key findings", styles["Heading3"]))
    for fnd in findings[:15]:
        story.append(Paragraph(f"• {fnd['text']}", styles["Normal"]))
        srcs = ", ".join(s for s in fnd.get("sources", []) if s)
        if srcs:
            story.append(Paragraph(f"   sources: {srcs}", styles["Italic"]))
    pdf.build(story)
    return out_path

def export_xlsx(doc, doc_records, comp, out_path):
    import openpyxl
    wb = openpyxl.Workbook()

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, min(60, len(str(h)) * 2))
        from openpyxl.styles import Font, PatternFill
        fill = PatternFill("solid", fgColor="1E3A8A")
        for c in ws[1]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = fill
        return ws

    ws = wb.active
    ws.title = "Overview"
    ws.append(["File", doc.get("filename")])
    ws.append(["Uploaded", time.strftime("%Y-%m-%d %H:%M")])
    ws.append(["Records extracted", len(doc_records)])
    ws.append(["Matched", len(comp["matched"])])
    ws.append(["Discrepancies", len(comp["discrepancies"])])
    ws.append(["New in document", len(comp["new_in_document"])])
    ws.append(["Not in document", len(comp["missing_in_document"])])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    sheet("A_Match", ["metric", "region", "period", "unit", "doc", "site", "abs_diff", "pct_diff", "doc_source", "site_source"],
          [[r["metric"], r["region"], r["period"], r["unit"], r["doc_value"], r["site_value"],
            r["abs_diff"], r["pct_diff"], r.get("doc_source"), r.get("site_source")] for r in comp["matched"]])
    sheet("B_Discrepancy", ["metric", "region", "period", "unit", "doc", "site", "pct_diff", "reasons", "doc_source", "site_source"],
          [[r["metric"], r["region"], r["period"], r["unit"], r["doc_value"], r["site_value"],
            r["pct_diff"], "/".join(r.get("possible_reasons", [])), r.get("doc_source"), r.get("site_source")]
           for r in comp["discrepancies"]])
    sheet("C_New", ["metric", "region", "period", "unit", "doc_value", "source"],
          [[r["metric"], r["region"], r["period"], r["unit"], r["doc_value"], r.get("doc_source")] for r in comp["new_in_document"]])
    sheet("D_Missing", ["metric", "region", "period", "unit", "site_value", "source"],
          [[r["metric"], r["region"], r["period"], r["unit"], r["site_value"], r.get("site_source")] for r in comp["missing_in_document"]])
    sheet("Extracted", ["metric", "region", "year", "period", "value", "unit", "sheet", "confidence"],
          [[r["metric_name"], r["region"], r["year"], r["period"], r["value"], r["unit"], r["source_sheet"], r["confidence"]]
           for r in doc_records])
    wb.save(out_path)
    return out_path
