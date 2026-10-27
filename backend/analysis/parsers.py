# -*- coding: utf-8 -*-
"""Generic document parsers: XLSX, CSV, TXT, DOCX, PDF.
Each returns:
  - text: full text (for search/OCR fallback)
  - tables: list of tables, each = {"sheet"/"page": label, "rows": [[...], ...]}
  - warnings: list of strings (e.g. 'no text layer, OCR needed')
"""
import os
import csv
import io
import re

MAX_SIZE = 25 * 1024 * 1024  # 25 MB

ALLOWED_EXT = {".pdf", ".docx", ".xlsx", ".csv", ".txt"}

DANGEROUS_PATTERNS = [
    b"\x7fELF",           # ELF binary
    b"PK\x05\x06",        # empty zip marker misuse
]

def validate_file(path, filename):
    """Return (ok, error_code, error_message)."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXT:
        return False, "bad_format", ext
    size = os.path.getsize(path)
    if size == 0:
        return False, "empty", filename
    if size > MAX_SIZE:
        return False, "too_large", f"{size} bytes"
    with open(path, "rb") as f:
        head = f.read(512)
    if ext in (".pdf",):
        if not head.startswith(b"%PDF"):
            return False, "corrupt", "not a valid PDF"
    if ext in (".xlsx",):
        if head[:4] != b"PK\x03\x04":
            return False, "corrupt", "not a valid XLSX (zip)"
    return True, None, None

def parse_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    tables = []
    text_parts = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append([("" if c is None else c) for c in r])
        if rows:
            tables.append({"sheet": ws.title, "rows": rows})
            for r in rows:
                line = " | ".join(str(c) for c in r if c != "")
                if line.strip():
                    text_parts.append(line)
    return {"text": "\n".join(text_parts), "tables": tables, "warnings": []}

def parse_csv(path):
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        raw = f.read()
    # try common delimiters
    sample = raw[:4000]
    delim = "\t" if sample.count("\t") > sample.count(";") > sample.count(",") else (
        ";" if sample.count(";") > sample.count(",") else ",")
    reader = csv.reader(io.StringIO(raw), delimiter=delim)
    rows = [row for row in reader]
    tables = [{"sheet": "CSV", "rows": rows}]
    text = "\n".join(" | ".join(r) for r in rows)
    return {"text": text, "tables": tables, "warnings": []}

def parse_txt(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return {"text": text, "tables": [], "warnings": []}

def parse_docx(path):
    import docx
    doc = docx.Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    tables = []
    for i, t in enumerate(doc.tables):
        rows = [[c.text for c in row.cells] for row in t.rows]
        tables.append({"sheet": f"Таблица {i+1}", "rows": rows})
        parts.append("\n".join(" | ".join(r) for r in rows))
    return {"text": "\n".join(parts), "tables": tables, "warnings": []}

def parse_pdf(path):
    warnings = []
    tables = []
    text_parts = []
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 200:
                warnings.append("Документ больше 200 страниц, обработаны первые 200")
            for i, page in enumerate(pdf.pages[:200], 1):
                txt = page.extract_text() or ""
                if txt.strip():
                    text_parts.append(f"[стр. {i}] {txt}")
                else:
                    warnings.append(f"Стр. {i}: текст не извлечён — возможно, сканер (нужен OCR)")
                try:
                    found = page.extract_tables()
                    for j, t in enumerate(found or [], 1):
                        rows = [[("" if c is None else str(c)) for c in row] for row in t]
                        tables.append({"sheet": f"стр. {i}, таблица {j}", "rows": rows})
                except Exception:
                    pass
    except Exception as e:
        # fallback to pypdf text-only
        warnings.append(f"pdfplumber не смог открыть файл ({e}), использован базовый извлекатель")
        try:
            from pypdf import PdfReader
            rd = PdfReader(path)
            for i, p in enumerate(rd.pages[:200], 1):
                t = p.extract_text() or ""
                if t.strip():
                    text_parts.append(f"[стр. {i}] {t}")
                else:
                    warnings.append(f"Стр. {i}: текст не извлечён — возможно, сканер (нужен OCR)")
        except Exception as e2:
            warnings.append(f"PDF повреждён или не читается: {e2}")
    return {"text": "\n".join(text_parts), "tables": tables, "warnings": warnings}

def parse_any(path, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".xlsx":
        return parse_xlsx(path)
    if ext == ".csv":
        return parse_csv(path)
    if ext in (".txt",):
        return parse_txt(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"unsupported: {ext}")
